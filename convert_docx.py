import os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
import markdown
from lxml import html

def add_formatted_text(paragraph, element):
    """Recursively add text with formatting (bold, italic) to a paragraph."""
    if element.text:
        run = paragraph.add_run(element.text)
        if element.tag == 'strong' or element.tag == 'b':
            run.bold = True
        if element.tag == 'em' or element.tag == 'i':
            run.italic = True

    for child in element:
        add_formatted_text(paragraph, child)
        if child.tail:
            paragraph.add_run(child.tail)

def md_to_docx(md_file, docx_file):
    with open(md_file, 'r', encoding='utf-8') as f:
        md_text = f.read()

    # Use extra features for better HTML
    html_text = markdown.markdown(md_text, extensions=['extra'])
    tree = html.fromstring(f"<div>{html_text}</div>")

    doc = Document()

    for element in tree:
        if element.tag in ['h1', 'h2', 'h3', 'h4']:
            level = int(element.tag[1])
            p = doc.add_heading('', level=level)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            add_formatted_text(p, element)
        elif element.tag == 'p':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            add_formatted_text(p, element)
        elif element.tag == 'ul' or element.tag == 'ol':
            for li in element:
                p = doc.add_paragraph(style='List Bullet' if element.tag == 'ul' else 'List Number')
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                add_formatted_text(p, li)

    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.rtl = True
        for run in paragraph.runs:
            run.font.size = Pt(12)
            run.font.name = 'Arial'

    doc.save(docx_file)

if __name__ == "__main__":
    md_to_docx('duga-book/book.md', 'duga-book/book.docx')
    print("DOCX generated with formatting.")
